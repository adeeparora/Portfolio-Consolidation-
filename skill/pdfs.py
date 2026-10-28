import sys
from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors

# usage: python3 pdfs.py <book.xlsx> <as-on dd.mm.yyyy> <close dd.mm.yyyy>
BOOK, ASON, CLOSE = sys.argv[1], sys.argv[2], sys.argv[3]
# Adeep's manual colour choice (same list as GREEN_NAMES in consolidate.py). Change only when he says so.
GREEN_NAMES = {
    'AFCOM Holdings Ltd', 'Affle 3i Ltd', 'Avenue Supermarts (DMart) Ltd', 'Bajaj Finance Ltd',
    'Bajaj Housing Finance Ltd', 'Connplex Cinemas Ltd', 'Eicher Motors Ltd', 'Fineotex Chemical Ltd',
    'HDFC Bank Ltd', 'HDFC Gold ETF', 'L&T Technology Services Ltd', 'LTM Ltd (formerly LTIMindtree)',
    'Nestle India Ltd', 'Nippon India ETF Gold BeES', 'Persistent Systems Ltd', 'Poonawalla Fincorp Ltd',
    'Rategain Travel Technologies Ltd', 'Tata Consultancy Services Ltd', 'Tata Elxsi Ltd',
    'Titan Company Ltd', 'Trent Ltd', 'Vinati Organics Ltd', 'Waaree Energies Ltd'}
wb = load_workbook(BOOK, data_only=True)
TABS = [n for n in wb.sheetnames if n not in ('Summary', 'Side by Side', 'Prices', '_state')]


def ab(b):   # shorter broker labels so multi-demat rows fit
    return (b.replace('Axis Direct', 'Axis').replace('HDFC Securities', 'HDFC')
             .replace('(Shilpa)', '(S)').replace('(Manish)', '(M)'))


def pc(v):
    return f"({abs(v)*100:,.0f}%)" if v < 0 else f"{v*100:,.0f}%"


for tab in TABS:
    ws = wb[tab]
    name, broker = ws['A1'].value.split(' – ')[:2]
    rows = []
    for r in range(7, 400):
        if not ws.cell(r, 2).value or ws.cell(r, 2).value == 'TOTAL':
            break
        rows.append((ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 5).value,
                     ws.cell(r, 7).value, ws.cell(r, 8).value))
    known = sorted([x for x in rows if x[3] is not None], key=lambda x: -(x[5] or 0))
    unk = [x for x in rows if x[3] is None]
    tot_mv = sum(x[5] or 0 for x in rows)
    tot_inv = sum(x[3] for x in known)
    mv_known = sum(x[5] or 0 for x in known)
    data = [["Rank", "Company", "Demat / Broker", "No. of\nShares", "Wtd Avg Cost\n(Rs.)", "Total Invested\n(Rs.)",
             "Market Price\n(Rs.)", "Market Value\n(Rs.)", "Unrealised P&L\n(Rs.)", "P&L %", "Portfolio\nWeight %"]]
    i = 0
    for sec, br, q, inv, px, mv in known:
        i += 1
        pl = mv - inv
        data.append([i, sec, ab(br), f"{q:,.0f}", (f"{inv/q:,.0f}" if inv else "0"), f"{inv:,.0f}",
                     f"{px:,.0f}" if px >= 100 else f"{px:,.2f}", f"{mv:,.0f}", f"{pl:,.0f}",
                     (pc(pl / inv) if inv else "N/A"), f"{mv/tot_mv*100:,.0f}%"])
    for sec, br, q, inv, px, mv in unk:
        i += 1
        data.append([i, sec + "*", ab(br), f"{q:,.0f}"] + ["N/A"] * 7)
    nrows = len(data)
    data.append(["", "Subtotal – holdings with known cost basis", "", "", "", f"{tot_inv:,.2f}", "", f"{mv_known:,.2f}", "", "", ""])
    data.append(["", "TOTAL MARKET VALUE – ALL HOLDINGS", "", "", "", "", "", f"{tot_mv:,.2f}", "", "", "100.0%"])
    t = Table(data, colWidths=[28, 175, 70, 42, 50, 62, 50, 62, 62, 38, 44], repeatRows=1)
    fs = 6.2 if len(rows) <= 33 else 4.8
    pd = 1.0 if len(rows) <= 33 else 0.1
    st = [('LEADING', (0, 1), (-1, -1), fs + 1.2), ('FONTSIZE', (0, 0), (-1, -1), fs),
          ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F3864')),
          ('TEXTCOLOR', (0, 0), (-1, 0), colors.white), ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
          ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('ALIGN', (3, 1), (-1, -1), 'RIGHT'), ('ALIGN', (0, 1), (0, -1), 'CENTER'),
          ('GRID', (0, 0), (-1, nrows - 1), 0.3, colors.grey), ('FONTNAME', (0, nrows), (-1, -1), 'Helvetica-Bold'),
          ('BACKGROUND', (0, nrows), (-1, -1), colors.HexColor('#DCE6F1')), ('BOX', (0, nrows), (-1, -1), 0.3, colors.grey),
          ('TOPPADDING', (0, nrows), (-1, -1), 4), ('BOTTOMPADDING', (0, nrows), (-1, -1), 4),
          ('TOPPADDING', (0, 0), (-1, nrows - 1), pd), ('BOTTOMPADDING', (0, 0), (-1, nrows - 1), pd)]
    for ri, rw in enumerate(data[1:nrows], 1):
        if rw[1].rstrip('*') in GREEN_NAMES:
            st.append(('TEXTCOLOR', (1, ri), (1, ri), colors.HexColor('#00B050')))
    t.setStyle(TableStyle(st))
    h = ParagraphStyle('h', fontName='Helvetica-Bold', fontSize=9, spaceAfter=8)
    sub = ParagraphStyle('s', fontName='Helvetica-Oblique', fontSize=6.5, spaceBefore=6, textColor=colors.grey)
    foot = Paragraph(f"Market prices are NSE/BSE closes of {CLOSE} (Screener.in); ETFs, rights entitlements and holdings marked * "
                     "use statement prices or carry no price. * No cost or price on the statement.", sub)
    fn = f"Consolidated_Portfolio_Summary_{name.replace(' ', '_')}_{CLOSE}_close.pdf"
    d = SimpleDocTemplate(fn, pagesize=landscape(A4), leftMargin=36, rightMargin=36, topMargin=22, bottomMargin=16)
    d.build([Paragraph(f"CONSOLIDATED PORTFOLIO SUMMARY – {name.upper()} ({broker.upper()}) – as on {ASON} (priced at {CLOSE} close)", h), t, foot])
    print(fn, len(rows), round(tot_mv), round(tot_inv))
